"""
tools/compile_results.py

Turns the evaluation's records into the numbers for the manuscript: the app's session logs (every laptop's
data/trainer_sessions folder) plus the Google Sheet "Evaluation records" downloaded as .xlsx
(File > Download > Microsoft Excel).

    python tools/compile_results.py --sheet Evaluation_Records.xlsx
    python tools/compile_results.py --sheet records.xlsx --root data/trainer_sessions --root E:/laptop2/trainer_sessions

The counting rules are the ones fixed in the procedure before any data was seen (docs/HANDOFF.md, Evaluation):
  - Steps: P4 = the first Practice; P5 = the first 10-repetition Drill after it that finished all 10 attempts;
    P8 = the first finished Challenge after P5; P9 = the first Combo; P10 = the first Match simulation.
  - Recognition rate per signal uses P5 + P8. Recognised = the model's detected signal is the prompted one.
    NO READING attempts (only in session_log.csv) count as misses and are also reported on their own.
  - Reported twice: all attempts, and with exclusions = Attempt notes rows on P5 or P8 typed Wrong signal,
    Out of frame or Other (matched by the on-screen attempt number).
  - Consistency: P5 attempts 1-5 give the same verdict at least 4 times.
  - Feedback within 3 s: P5 attempts' feedback_latency_s.
  - Whistle detection: Whistles tab, (flashes - flashes with no whistle) / blown, cross-checked against the logs.
  - Only codes starting G- are participants; DEMO, DRY-01 and every other folder are left out.
Writes data/results/results.md plus CSV tables, and lists every problem it finds (missing sessions, settings that
differ from the procedure, sheet and logs that disagree) so they can be fixed before the numbers are used.
Run it from the project root.
"""

import argparse
import csv
import glob
import json
import os
import re
import sys
from collections import Counter, defaultdict

SIGNALS = ["team_to_serve_left", "team_to_serve_right", "service_authorization_left", "service_authorization_right",
           "ball_out", "ball_in", "double_contact", "end_of_set"]
SHEET_SIGNAL = {   # the Drill signal names used in the records sheet
    "team to serve – left arm": "team_to_serve_left", "team to serve – right arm": "team_to_serve_right",
    "service authorization – left arm": "service_authorization_left",
    "service authorization – right arm": "service_authorization_right",
    "ball out": "ball_out", "ball in": "ball_in", "double contact": "double_contact", "end of set": "end_of_set"}
EXCLUDING_TYPES = {"wrong signal", "out of frame", "other"}
FOLDER_RE = re.compile(r"^(\d{8}_\d{6})_(practice|drill|challenge|combo|sim|match_test)(_correct|_wrong)?$")


def pct(a, b):
    return f"{100.0 * a / b:.1f}%" if b else "-"


def read_csv(path):
    if not os.path.exists(path):
        return []
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# ---------------------------------------------------------------------------------------------------------------
# The records sheet
# ---------------------------------------------------------------------------------------------------------------

def read_sheet(path):
    """{tab name: [row dict keyed by header]}; the grey EXAMPLE row and empty rows are skipped."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        head_i = next((i for i, r in enumerate(rows[:10]) if r and str(r[0] or "").strip() in ("Code", "Step", "Attempt")),
                      None)
        if head_i is None:
            continue
        head = [str(h or "").strip() for h in rows[head_i]]
        data = []
        for r in rows[head_i + 1:]:
            if not r or r[0] is None or str(r[0]).strip() in ("", "EXAMPLE"):
                continue
            data.append({h: v for h, v in zip(head, r) if h})
        out[ws.title.strip()] = data
    return out


def col(row, prefix):
    """A sheet cell by the start of its header ("Dataset volunteer" matches "Dataset volunteer (Y/N)")."""
    p = prefix.lower()
    for k, v in row.items():
        if k.lower().startswith(p):
            return "" if v is None else (str(v).strip() if not isinstance(v, (int, float)) else v)
    return ""


def num(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------------------------------------------
# The session logs
# ---------------------------------------------------------------------------------------------------------------

def read_session(path):
    m = FOLDER_RE.match(os.path.basename(path))
    s = {"path": path, "name": os.path.basename(path), "mode": m.group(2), "intent": (m.group(3) or "_normal")[1:]}
    p = os.path.join(path, "summary.json")
    try:
        s["summary"] = json.load(open(p, encoding="utf-8")) if os.path.exists(p) else None
    except (OSError, ValueError):
        s["summary"] = None
    s["log"] = read_csv(os.path.join(path, "session_log.csv"))
    start = next((r.get("detail", "") for r in s["log"] if r.get("event") == "session_start"), "")
    s["settings"] = dict(kv.split("=", 1) for kv in (x.strip() for x in start.split(";")) if "=" in kv)
    s["reps"] = num(s["settings"].get("repetitions"))
    s["rows"] = read_csv(os.path.join(path, "attempts.csv"))
    s["whistles"] = read_csv(os.path.join(path, "whistle_events.csv"))
    summ = s["summary"] or {}
    s["problem"] = ("no summary.json (the app stopped before the session ended)" if s["summary"] is None
                    else f"error: {summ['error']}" if summ.get("error")
                    else "camera lost" if summ.get("camera_lost") else "")
    return s


def attempt_units(s):
    """One unit per graded attempt of a session, in order: the counted rows of attempts.csv plus the NO READING
    attempts that only session_log.csv records. Each: step (on-screen number), target, detected, verdict,
    latency, no_reading."""
    units = []
    rows = [r for r in s["rows"] if r.get("kind") == "gesture"]
    has_step = rows and "step_no" in rows[0]
    for i, r in enumerate(rows):
        units.append({"step": num(r.get("step_no")) if has_step else i + 1, "target": r["target"],
                      "detected": r.get("detected", ""), "verdict": r["verdict"],
                      "latency": float(r["feedback_latency_s"]) if r.get("feedback_latency_s") else None,
                      "no_reading": False, "t": r.get("ph_time", "")})
    for e in s["log"]:
        if e.get("event") != "verdict" or "NO_READING" not in e.get("detail", ""):
            continue
        m = re.search(r"; step (\d+)", e["detail"])
        for part in e["detail"].split(";"):
            lab, _, rest = part.strip().partition("=")
            if lab in SIGNALS and rest.startswith("NO_READING"):
                units.append({"step": int(m.group(1)) if m else None, "target": lab, "detected": "", "verdict":
                              "NO_READING", "latency": None, "no_reading": True, "t": e.get("ph_time", "")})
    units.sort(key=lambda u: u["t"])
    return units


def find_steps(sessions, problems, code):
    """Which session folder is which step of the procedure."""
    normal = [s for s in sessions if s["intent"] == "normal"]
    steps = {}
    practice = [s for s in normal if s["mode"] == "practice"]
    if practice:
        steps["P4"] = practice[0]
    else:
        problems.append(f"{code}: no Practice session (P4); P5 taken as the first 10-repetition Drill")
    after = steps["P4"]["name"] if "P4" in steps else ""
    for s in normal:
        if s["mode"] != "drill" or s["name"] < after or s["reps"] != 10:
            continue
        n_steps = len({u["step"] for u in attempt_units(s) if not u["no_reading"]})
        if s["problem"] or n_steps < 10:
            problems.append(f"{code}: 10-repetition Drill {s['name']} did not finish ({s['problem'] or f'{n_steps}/10 attempts'}); not used as P5")
            continue
        steps["P5"] = s
        break
    if "P5" not in steps:
        problems.append(f"{code}: no finished 10-repetition Drill (P5)")
    p5 = steps["P5"]["name"] if "P5" in steps else after
    for key, mode in (("P8", "challenge"), ("P9", "combo"), ("P10", "sim")):
        cands = [s for s in normal if s["mode"] == mode and s["name"] > p5]
        done = [s for s in cands if not s["problem"]]
        if cands and not done:
            problems.append(f"{code}: every {mode} session stopped early ({', '.join(s['name'] for s in cands)})")
        if done:
            steps[key] = done[0]
            for s in cands:
                if s is done[0]:
                    break
                problems.append(f"{code}: {mode} session {s['name']} stopped early ({s['problem']}); the next one is used as {key}")
        elif not cands:
            problems.append(f"{code}: no {mode} session ({key})")
    return steps


def check_settings(code, steps, problems):
    want = {"P5": ("drill", False), "P8": ("challenge", False), "P9": ("combo", True), "P10": ("sim", True)}
    for key, (_mode, whistle) in want.items():
        s = steps.get(key)
        if not s:
            continue
        st = s["settings"]
        if st.get("level") and st["level"] != "standard":
            problems.append(f"{code}: {key} ({s['name']}) ran at {st['level']}, not Standard")
        w = st.get("whistle_option")
        if w is not None and (w == "True") != whistle:
            problems.append(f"{code}: {key} ({s['name']}) had Require the whistle {'on' if w == 'True' else 'off'}; "
                            f"the procedure says {'on' if whistle else 'off'}")
    s = steps.get("P10")
    if s and s["settings"].get("continuous") == "True":
        problems.append(f"{code}: P10 ran with Continuous on")


# ---------------------------------------------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sheet", required=True, help="the records sheet downloaded as .xlsx")
    ap.add_argument("--root", action="append", default=None,
                    help="a trainer_sessions folder; repeat for every laptop's copy (default data/trainer_sessions)")
    ap.add_argument("--prefix", default="G-", help="participant codes start with this (default G-)")
    ap.add_argument("--out", default=os.path.join("data", "results"))
    args = ap.parse_args()
    roots = args.root or [os.path.join("data", "trainer_sessions")]
    problems = []

    sheet = read_sheet(args.sheet)
    for tab in ("Participants", "Attempt notes", "Whistles", "Help log"):
        if tab not in sheet:
            sys.exit(f"The sheet has no '{tab}' tab: is this the Evaluation records v2.1 download?")
    people = {str(col(r, "Code")): r for r in sheet["Participants"] if str(col(r, "Code")).startswith(args.prefix)}

    # every laptop's folders, by code
    folders = defaultdict(list)
    for root in roots:
        for d in sorted(glob.glob(os.path.join(root, args.prefix + "*"))):
            if os.path.isdir(d):
                folders[os.path.basename(d)].append(d)
    sessions = {}
    for code, dirs in folders.items():
        if len(dirs) > 1:
            problems.append(f"{code}: has folders on more than one laptop ({'; '.join(dirs)}); codes must be unique, "
                            "check that these are the same person")
        sessions[code] = sorted((read_session(p) for d in dirs for p in glob.glob(os.path.join(d, "*"))
                                 if os.path.isdir(p) and FOLDER_RE.match(os.path.basename(p))), key=lambda s: s["name"])

    done = sorted(c for c, r in people.items() if col(r, "Setup"))
    codes = sorted(set(done) | set(sessions))
    for c in codes:
        if c not in sessions:
            problems.append(f"{c}: Setup is filled in the sheet but no session folder was found (typo in the code, or "
                            "that laptop's folder not added with --root?)")
        elif c not in people or not col(people[c], "Setup"):
            problems.append(f"{c}: has session folders but its Participants row has no Setup (fill the sheet row)")
    codes = [c for c in codes if c in sessions]

    notes = defaultdict(list)
    for r in sheet["Attempt notes"]:
        notes[str(col(r, "Code"))].append(r)
    whistle_rows = defaultdict(dict)
    for r in sheet["Whistles"]:
        whistle_rows[str(col(r, "Code"))][str(col(r, "Step"))] = r

    units_out, per_person, pair_rows = [], [], []
    for code in codes:
        p = people.get(code, {})
        setup = col(p, "Setup") or "unknown"
        steps = find_steps(sessions[code], problems, code)
        check_settings(code, steps, problems)
        assigned = SHEET_SIGNAL.get(str(col(p, "Drill signal")).lower())
        person = {"code": code, "setup": setup, "location": col(p, "Location"), "researcher": col(p, "Researcher"),
                  "dataset_volunteer": col(p, "Dataset volunteer"), "crash_sheet": col(p, "Crash"),
                  "steps": " ".join(f"{k}={v['name']}" for k, v in sorted(steps.items(), key=lambda kv: int(kv[0][1:])))}
        for s in sessions[code]:
            if s["problem"]:
                problems.append(f"{code}: session {s['name']} ({s['mode']}): {s['problem']}")
        for key, s in steps.items():
            if s["rows"] and "detected" not in s["rows"][0]:
                problems.append(f"{code}: {key} ({s['name']}) was recorded by an older version of the app (no 'detected' "
                                "column), so its recognition cannot be counted; that laptop needs git pull")
        person["crash_logs"] = sum(1 for s in sessions[code] if s["problem"])

        # sheet completeness
        deliberate = [r for r in notes.get(code, []) if str(col(r, "Step")) == "P7" and str(col(r, "Type")).lower() == "deliberate"]
        if len(deliberate) != 2:
            problems.append(f"{code}: Attempt notes has {len(deliberate)} P7 Deliberate rows (expected 2)")
        for st in ("P9", "P10"):
            wr = whistle_rows.get(code, {}).get(st)
            if wr is None or num(col(wr, "Whistles blown")) is None:
                problems.append(f"{code}: Whistles {st} row is empty")
        if not col(p, "Researcher"):
            problems.append(f"{code}: Participants row has no Researcher")

        # recognition units (P5 + P8) with the exclusions from Attempt notes
        excl = defaultdict(list)
        for r in notes.get(code, []):
            st, typ, k = str(col(r, "Step")), str(col(r, "Type")).lower(), num(col(r, "Attempt #"))
            if st in ("P5", "P8") and typ in EXCLUDING_TYPES:
                if k is None:
                    problems.append(f"{code}: Attempt notes row ({st}, {typ}) has no Attempt #; it cannot be matched")
                else:
                    excl[(st, k)].append(typ)
        for key in ("P5", "P8"):
            s = steps.get(key)
            if not s:
                continue
            us = attempt_units(s)
            if key == "P5" and assigned and us and {u["target"] for u in us} != {assigned}:
                problems.append(f"{code}: P5 drilled {sorted({u['target'] for u in us})}, but the sheet assigns {assigned}")
            used = set()
            for (st, k), types in excl.items():
                if st != key:
                    continue
                cand = [u for u in us if u["step"] == k and id(u) not in used]
                if not cand:
                    problems.append(f"{code}: Attempt notes {st} attempt {k} matches no attempt in {s['name']}")
                    continue
                pick = (next((u for u in cand if u["no_reading"]), cand[0]) if "out of frame" in types
                        else next((u for u in cand if not u["no_reading"]), cand[0]))
                used.add(id(pick))
                pick["excluded"] = ", ".join(types)
            for u in us:
                units_out.append({"code": code, "setup": setup, "step": key, "session": s["name"], "attempt": u["step"],
                                  "target": u["target"], "detected": u["detected"], "verdict": u["verdict"],
                                  "recognised": int(u["detected"] == u["target"]), "no_reading": int(u["no_reading"]),
                                  "excluded": u.get("excluded", ""),
                                  "latency_s": "" if u["latency"] is None else u["latency"]})

        # consistency and feedback time (P5)
        s5 = steps.get("P5")
        first = {}
        for u in (attempt_units(s5) if s5 else []):
            if not u["no_reading"] and u["step"] not in first:
                first[u["step"]] = u
        five = [first[k]["verdict"] for k in range(1, 6) if k in first]
        person["p5_first5"] = " ".join(v[0] for v in five)
        person["consistent"] = (len(five) == 5 and Counter(five).most_common(1)[0][1] >= 4) if five else ""
        lat = [u["latency"] for u in first.values() if u["latency"] is not None]
        person["fb_n"], person["fb_within3"] = len(lat), sum(1 for x in lat if x <= 3.0)
        person["fb_max"] = max(lat) if lat else ""

        # P9 / P10: each signal of the pair, and the whistle per attempt
        for key in ("P9", "P10"):
            s = steps.get(key)
            if not s:
                continue
            for r in s["rows"]:
                pair_rows.append({"code": code, "step": key, "kind": r["kind"], "target": r["target"],
                                  "detected": r.get("detected", ""), "verdict": r["verdict"]})
            auto = sum(1 for w in s["whistles"] if w.get("source") == "auto")
            manual = sum(1 for w in s["whistles"] if w.get("source") == "manual")
            wr = whistle_rows.get(code, {}).get(key)
            flashes = num(col(wr, "WHISTLE! flashes")) if wr else None
            if manual:
                problems.append(f"{code}: {key} used the W key {manual} time(s) instead of a detected whistle")
            if flashes is not None and flashes != auto:
                problems.append(f"{code}: {key} sheet says {flashes} WHISTLE! flashes, the log recorded {auto} detections")
            person[f"{key}_log_detections"] = auto
        per_person.append(person)

    # ---- totals ----
    os.makedirs(args.out, exist_ok=True)

    def write(name, rows):
        if rows:
            fields = list(dict.fromkeys(k for r in rows for k in r))    # rows can differ (a missing step adds no key)
            with open(os.path.join(args.out, name), "w", newline="", encoding="utf-8") as f:
                w = csv.DictWriter(f, fieldnames=fields)
                w.writeheader()
                w.writerows(rows)

    write("recognition_attempts.csv", units_out)
    write("participants.csv", per_person)
    write("pair_attempts.csv", pair_rows)

    lines = ["# Evaluation results (compiled)", "",
             f"Participants: {len(codes)} ({', '.join(codes) or 'none'}). Setups: "
             + ", ".join(f"{k} {v}" for k, v in Counter(p['setup'] for p in per_person).items()) + ".",
             f"Sources: {', '.join(roots)} and {os.path.basename(args.sheet)}.", ""]

    def rec_table(title, rows):
        lines.extend([f"## {title}", "",
                      "| Signal | Attempts | Recognised | Rate | NO READING | Excluded | Attempts after exclusions | Recognised | Rate |",
                      "|---|---|---|---|---|---|---|---|---|"])
        tot = Counter()
        for lab in SIGNALS + ["all"]:
            rs = rows if lab == "all" else [u for u in rows if u["target"] == lab]
            kept = [u for u in rs if not u["excluded"]]
            c = Counter(n=len(rs), rec=sum(u["recognised"] for u in rs), nr=sum(u["no_reading"] for u in rs),
                        ex=len(rs) - len(kept), kn=len(kept), krec=sum(u["recognised"] for u in kept))
            if lab != "all" and not c["n"]:
                continue
            name = "**All signals**" if lab == "all" else lab
            lines.append(f"| {name} | {c['n']} | {c['rec']} | {pct(c['rec'], c['n'])} | {c['nr']} | {c['ex']} | "
                         f"{c['kn']} | {c['krec']} | {pct(c['krec'], c['kn'])} |")
        lines.append("")

    rec_table("Recognition rate per signal (P5 + P8)", units_out)
    for setup in sorted({u["setup"] for u in units_out}):
        rec_table(f"Recognition, {setup} sessions only", [u for u in units_out if u["setup"] == setup])
    for key in ("P5", "P8"):
        rec_table(f"Recognition, {key} only", [u for u in units_out if u["step"] == key])

    v = Counter(u["verdict"] for u in units_out if u["step"] == "P5")
    lines += ["## P5 verdicts at Standard", "", ", ".join(f"{k} {n}" for k, n in v.most_common()) or "-", ""]

    cons = [p for p in per_person if p["consistent"] != ""]
    fb_n, fb_ok = sum(p["fb_n"] for p in per_person), sum(p["fb_within3"] for p in per_person)
    fb_max = max([p["fb_max"] for p in per_person if p["fb_max"] != ""], default="-")
    lines += ["## Consistency and feedback time (P5)", "",
              f"- Same verdict at least 4 of the first 5 attempts: {sum(1 for p in cons if p['consistent'])}/{len(cons)} participants.",
              f"- Verdict shown within 3 s of the arms coming down: {fb_ok}/{fb_n} attempts ({pct(fb_ok, fb_n)}); slowest {fb_max} s.",
              ""]

    lines += ["## Combo (P9) and Match simulation (P10), reported separately", "",
              "| Step | Signal attempts | Recognised | Rate | Whistle heard / attempts |", "|---|---|---|---|---|"]
    for key in ("P9", "P10"):
        g = [r for r in pair_rows if r["step"] == key and r["kind"] == "gesture"]
        w = [r for r in pair_rows if r["step"] == key and r["kind"] == "whistle"]
        rec = sum(1 for r in g if r["detected"] == r["target"])
        lines.append(f"| {key} | {len(g)} | {rec} | {pct(rec, len(g))} | "
                     f"{sum(1 for r in w if r['detected'] == 'ok')}/{len(w)} |")
    lines.append("")

    wl = [r for c in codes for r in whistle_rows.get(c, {}).values()]
    def wsum(rows, name):
        return sum(num(col(r, name)) or 0 for r in rows)
    lines += ["## Whistle detection (Whistles tab)", "", "| Part | Blown | Flashes | Flashes with no whistle | Detected | Detection rate |",
              "|---|---|---|---|---|---|"]
    for label, rows in (("All participants", wl), ("P9", [r for r in wl if str(col(r, "Step")) == "P9"]),
                        ("P10", [r for r in wl if str(col(r, "Step")) == "P10"])):
        b, f, fa = wsum(rows, "Whistles blown"), wsum(rows, "WHISTLE! flashes"), wsum(rows, "Flashes with no whistle")
        lines.append(f"| {label} | {b} | {f} | {fa} | {f - fa} | {pct(f - fa, b)} |")
    demo = {str(col(r, "Step"))[:6]: r for r in sheet["Whistles"] if str(col(r, "Code")) == "DEMO"}
    if demo:
        d, e = demo.get("D4 (d)"), demo.get("D4 (e)")
        lines.append("")
        lines.append(f"Demonstration: 5 whistles gave {(col(d, 'WHISTLE! flashes') if d else '') or 'not filled'} flashes; "
                     f"the no-whistle minute gave {(col(e, 'Flashes with no whistle') if e else '') or 'not filled'} false alarms.")
    lines.append("")

    help_rows = [r for r in sheet["Help log"] if str(col(r, "Code")) in codes]
    by_step = Counter((str(col(r, "Step")), str(col(r, "Kind"))) for r in help_rows)
    lines += ["## Help and steps redone", ""]
    lines += [f"- {st}: {by_step[(st, 'Help given')]} help, {by_step[(st, 'Step redone')]} redone"
              for st in sorted({k[0] for k in by_step}, key=lambda x: int(x[1:]) if x[1:].isdigit() else 99)] or ["- none"]
    lines += ["", "## Crashes and freezes", "",
              f"- Sheet (Crash or freeze = Yes): {sum(1 for p in per_person if str(p['crash_sheet']).lower() == 'yes')} participants.",
              f"- Logs (session stopped by an error, a lost camera, or with no summary): "
              f"{sum(p['crash_logs'] for p in per_person)} sessions.", ""]
    vol = [p["code"] for p in per_person if str(p["dataset_volunteer"]).lower() in ("yes", "y")]
    lines += [f"Dataset volunteers among participants: {len(vol)} ({', '.join(vol) or 'none'}).", ""]
    lines += ["## Problems to fix before using these numbers", ""] + ([f"- {x}" for x in problems] or ["- none"])
    with open(os.path.join(args.out, "results.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print("\n".join(lines))
    print(f"\nSaved {args.out}{os.sep}results.md, recognition_attempts.csv, participants.csv, pair_attempts.csv")


if __name__ == "__main__":
    main()
